"""Отчёты, фильтры, экспорт и веб-интерфейс."""

from __future__ import annotations

from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.constants import ThreadStatus
from app.models import Thread
from app.services.analysis_service import analyze_threads
from app.services.demo_data import load_demo_data
from app.services.export_excel import build_workbook
from app.services.reports import (
    ThreadFilters,
    daily_report,
    dashboard_metrics,
    day_bounds_utc,
    manager_statistics,
    query_threads,
    search_variants,
    sort_threads,
)
from app.services.working_time import WorkSchedule

SCHEDULE = WorkSchedule()


@pytest.fixture
def loaded(db):
    load_demo_data(db)
    analyze_threads(db, force=True, use_ai=False)
    return db


def data_day(db):
    """День (по местному времени), на который приходится большинство демо-обращений."""
    from collections import Counter

    from app.services.reports import day_bounds_utc

    threads = db.scalars(select(Thread)).all()
    days = Counter()
    for thread in threads:
        if thread.last_message_at is None:
            continue
        for offset in (0, 1):
            candidate = (thread.last_message_at + timedelta(days=offset)).date()
            start, end = day_bounds_utc(candidate, SCHEDULE)
            if start <= thread.last_message_at < end:
                days[candidate] += 1
                break
    return days.most_common(1)[0][0] if days else date.today()


def test_day_bounds_convert_timezone():
    start, end = day_bounds_utc(date(2026, 7, 21), SCHEDULE)
    # Москва UTC+3: локальные сутки начинаются в 21:00 UTC предыдущего дня
    assert start == datetime(2026, 7, 20, 21, 0)
    assert end == datetime(2026, 7, 21, 21, 0)


def test_dashboard_metrics(loaded):
    metrics = dashboard_metrics(loaded, date.today(), SCHEDULE)
    assert metrics["total_critical_all"] >= 3
    assert metrics["read_unanswered"] >= 1
    assert metrics["total_incomplete_all"] >= 5
    assert isinstance(metrics["incoming_messages"], int)


def test_daily_report_sections(loaded):
    report = daily_report(loaded, data_day(loaded))
    sections = report["sections"]
    for key in [
        "critical", "read_unanswered", "answered_incomplete", "partially_done",
        "waiting_client", "completed", "repeats", "longest_waiting", "needs_review",
    ]:
        assert key in sections
    assert report["categories"]
    assert isinstance(report["managers"], list)


def test_report_partially_done_has_mixed_tasks(loaded):
    report = daily_report(loaded, data_day(loaded))
    # Сценарий 12: часть задач выполнена, часть нет
    assert report["sections"]["partially_done"]


def test_filter_by_status(loaded):
    threads = query_threads(loaded, ThreadFilters(status=[ThreadStatus.CRITICAL]), SCHEDULE)
    assert threads
    assert all(t.current_status == ThreadStatus.CRITICAL for t in threads)


def test_filter_read_unanswered(loaded):
    threads = query_threads(loaded, ThreadFilters(read_unanswered=True), SCHEDULE)
    assert threads
    assert all(t.is_read_unanswered for t in threads)


def test_filter_by_request_type(loaded):
    threads = query_threads(loaded, ThreadFilters(request_type="invoice"), SCHEDULE)
    assert threads


def test_not_relevant_hidden_by_default(loaded):
    default = query_threads(loaded, ThreadFilters(), SCHEDULE)
    assert all(t.current_status != ThreadStatus.NOT_RELEVANT for t in default)
    with_spam = query_threads(loaded, ThreadFilters(include_not_relevant=True), SCHEDULE)
    assert len(with_spam) > len(default)


def test_search_finds_by_body_and_email(loaded):
    assert query_threads(loaded, ThreadFilters(search="подшипники"), SCHEDULE)
    assert query_threads(loaded, ThreadFilters(search="stroymash"), SCHEDULE)


def test_search_ignores_yo_difference(loaded):
    with_yo = query_threads(loaded, ThreadFilters(search="счёт"), SCHEDULE)
    without_yo = query_threads(loaded, ThreadFilters(search="счет"), SCHEDULE)
    assert with_yo and without_yo
    assert {t.id for t in with_yo} == {t.id for t in without_yo}


def test_search_variants():
    assert set(search_variants("счет")) == {"%счет%", "%счёт%"}


def test_sort_priority_puts_critical_first(loaded):
    threads = query_threads(loaded, ThreadFilters(sort="priority"), SCHEDULE)
    assert threads[0].current_status == ThreadStatus.CRITICAL


def test_sort_by_waiting(loaded):
    threads = sort_threads(loaded.scalars(select(Thread)).all(), "waiting")
    waits = [t.waiting_seconds or 0 for t in threads]
    assert waits == sorted(waits, reverse=True)


def test_manager_statistics(loaded):
    rows = manager_statistics(loaded, date.today() - timedelta(days=30), date.today())
    assert rows
    row = rows[0]
    for key in ["name", "threads", "completed", "incomplete", "no_reply", "completion_rate"]:
        assert key in row
    assert 0 <= row["completion_rate"] <= 100


def test_excel_export_structure(loaded):
    buffer = build_workbook(loaded, data_day(loaded))
    wb = load_workbook(buffer)
    assert wb.sheetnames == [
        "Сводка", "Критические", "Запрос не выполнен",
        "Все обращения", "По менеджерам", "По типам запросов",
    ]
    assert wb["Все обращения"].max_row > 1
    assert wb["Сводка"]["A1"].value.startswith("Аналитика")


# ------------------------------------------------------------------
# Веб-интерфейс
# ------------------------------------------------------------------

@pytest.fixture
def client(loaded, monkeypatch):
    from fastapi.testclient import TestClient

    import app.core.db as core_db
    from app.main import app

    session = loaded
    monkeypatch.setattr(core_db, "SessionLocal", lambda: session)

    def _get_db():
        yield session

    from app.core.db import get_db

    app.dependency_overrides[get_db] = _get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.mark.parametrize(
    "url",
    ["/", "/threads", "/report", "/managers", "/settings", "/demo", "/guide",
     "/setup", "/limitations", "/health"],
)
def test_pages_render(client, url):
    response = client.get(url)
    assert response.status_code == 200, url


def test_thread_detail_renders(client, loaded):
    thread = loaded.scalar(select(Thread))
    response = client.get(f"/threads/{thread.id}")
    assert response.status_code == 200
    assert "Запрос клиента" in response.text


def test_missing_thread_returns_404(client):
    assert client.get("/threads/999999").status_code == 404


def test_dashboard_shows_priority_sections(client):
    text = client.get("/").text
    assert "ПРОЧИТАНО, НО НЕ ОТВЕЧЕНО" in text
    assert "ОТВЕТ ОТПРАВЛЕН, НО ЗАПРОС КЛИЕНТА НЕ ВЫПОЛНЕН" in text


def test_html_export_downloads(client):
    response = client.get("/export/html")
    assert response.status_code == 200
    assert "attachment" in response.headers.get("content-disposition", "")
    assert "Аналитика клиентской почты" in response.text


def test_excel_export_downloads(client):
    response = client.get("/export/excel")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )


def test_manual_status_change_via_web(client, loaded):
    thread = loaded.scalar(select(Thread).where(Thread.current_status == ThreadStatus.CRITICAL))
    response = client.post(
        f"/threads/{thread.id}/status",
        data={"status": "COMPLETED", "comment": "Ответили по телефону"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    loaded.refresh(thread)
    assert thread.current_status == ThreadStatus.COMPLETED
    assert thread.manual_override is True
    assert thread.reviews


def test_sync_blocked_without_config(client):
    response = client.post("/actions/sync", data={"days": 7}, follow_redirects=False)
    assert response.status_code == 303
    assert "msg=" in response.headers["location"]


def test_dashboard_shows_open_state_regardless_of_day(loaded):
    """Показатели текущего состояния не должны обнуляться при выборе другой даты."""
    from datetime import date as _date

    far_future = _date(2030, 1, 1)
    metrics = dashboard_metrics(loaded, far_future, SCHEDULE)
    assert metrics["incoming_messages"] == 0          # за выбранный день писем нет
    assert metrics["total_critical_all"] >= 3         # но открытые обращения видны
    assert metrics["total_incomplete_all"] >= 5
    assert metrics["read_unanswered"] >= 1


# ------------------------------------------------------------------
# Демонстрационные данные должны быть явно обозначены
# ------------------------------------------------------------------

def test_demo_banner_visible_on_every_page(client):
    """Пользователь должен видеть, что перед ним тестовые письма, а не его почта."""
    for url in ("/", "/threads", "/report", "/managers"):
        text = client.get(url).text
        assert "Это демонстрационные данные, а не ваша почта" in text, url
        assert "demo@example.local" in text, url


def test_demo_rows_marked_in_table(client):
    assert 'class="chip demo"' in client.get("/threads").text


def test_demo_marker_in_thread_card(client, loaded):
    thread = loaded.scalar(select(Thread))
    text = client.get(f"/threads/{thread.id}").text
    assert "демонстрационное письмо, не из вашей почты" in text


def test_demo_data_can_be_removed(client, loaded):
    from app.models import MailAccount, Message

    assert loaded.scalars(select(Thread)).all()
    response = client.post("/actions/demo/clear", follow_redirects=False)
    assert response.status_code == 303

    demo = loaded.scalar(select(MailAccount).where(MailAccount.is_demo.is_(True)))
    assert demo is not None  # сам ящик остаётся, чтобы демо можно было загрузить снова
    assert loaded.scalars(select(Thread).where(Thread.account_id == demo.id)).all() == []
    assert loaded.scalars(select(Message).where(Message.account_id == demo.id)).all() == []
    assert "Это демонстрационные данные" not in client.get("/").text


def test_guide_page_covers_real_mail_setup(client):
    """Инструкция должна вести пользователя от пароля приложения до ежедневной работы."""
    text = client.get("/guide").text
    for fragment in [
        "Пароли приложений",
        "Настройки → Почтовые программы",
        "CORPORATE_DOMAINS",
        "Отправленные",
        "Проверить подключение",
        "Удалить демо-данные",
    ]:
        assert fragment in text, fragment


# ------------------------------------------------------------------
# Мастер подключения почты
# ------------------------------------------------------------------

def test_setup_page_renders(client):
    text = client.get("/setup").text
    assert "Подключение вашей почты" in text
    assert "Пароли приложений" in text
    assert "Домен вашей компании" in text


def test_setup_rejects_bad_email(client):
    response = client.post(
        "/actions/setup",
        data={"email": "не-адрес", "password": "x" * 16, "domains": "company.ru"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    from urllib.parse import unquote

    assert "kind=error" in response.headers["location"]
    assert "неверно" in unquote(response.headers["location"])


def test_setup_requires_domain(client):
    response = client.post(
        "/actions/setup",
        data={"email": "sales@company.ru", "password": "x" * 16, "domains": ""},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert "kind=error" in response.headers["location"]


def test_setup_saves_settings_without_touching_db_password(client, loaded, tmp_path, monkeypatch):
    """Пароль уходит в .env, а домен — в настройки. В базе пароля быть не должно."""
    import app.services.env_writer as env_writer
    from app.models import AppSetting

    env_file = tmp_path / ".env"
    env_file.write_text("YANDEX_EMAIL=\nYANDEX_APP_PASSWORD=\n", encoding="utf-8")
    monkeypatch.setattr(env_writer, "ENV_PATH", env_file)

    client.post(
        "/actions/setup",
        data={
            "email": "sales@mycompany.ru",
            "password": "abcdefghijklmnop",
            "domains": "mycompany.ru, mycompany.com",
        },
        follow_redirects=False,
    )

    saved = env_file.read_text(encoding="utf-8")
    assert "YANDEX_EMAIL=sales@mycompany.ru" in saved
    assert "abcdefghijklmnop" in saved

    rows = loaded.scalars(select(AppSetting)).all()
    assert all("abcdefghijklmnop" not in (r.value or "") for r in rows), (
        "Пароль не должен попадать в базу данных"
    )

    from app.services.settings_service import get_all_settings

    assert "mycompany.ru" in get_all_settings(loaded)["corporate.domains"]
